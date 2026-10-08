#!/usr/bin/env python3
"""Tao file BANG THEO DOI CONG VIEC TUAN tu anh chup."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else "BANG_THEO_DOI_CONG_VIEC_TUAN.xlsx"

THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

wb = Workbook()
ws = wb.active
ws.title = "TheoDoiCongViec"

NCOL = 7  # STT, Ngay, Cong viec, Nguoi thuc hien, So gio, Trang thai, Ghi chu

# 1) Tieu de
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NCOL)
c = ws.cell(row=1, column=1, value="BẢNG THEO DÕI CÔNG VIỆC TUẦN")
c.font = Font(bold=True, size=16)
c.alignment = CENTER
ws.row_dimensions[1].height = 30

# 2) Dong Bo phan / Tuan
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=2)
ws.cell(row=2, column=1, value="Bộ phận:").font = Font(bold=True)
ws.cell(row=2, column=1).alignment = LEFT
ws.merge_cells(start_row=2, start_column=3, end_row=2, end_column=4)
ws.cell(row=2, column=3).alignment = LEFT

ws.merge_cells(start_row=2, start_column=5, end_row=2, end_column=6)
ws.cell(row=2, column=5, value="Tuần: ...../...../2026").font = Font(bold=True)
ws.cell(row=2, column=5).alignment = LEFT
ws.merge_cells(start_row=2, start_column=7, end_row=2, end_column=7)
ws.cell(row=2, column=7).alignment = LEFT
ws.row_dimensions[2].height = 22

# 3) Header
headers = ["STT", "Ngày", "Công việc", "Người thực hiện", "Số giờ", "Trạng thái", "Ghi chú"]
HR = 3
for j, h in enumerate(headers, 1):
    c = ws.cell(row=HR, column=j, value=h)
    c.font = Font(bold=True)
    c.alignment = CENTER
    c.fill = HEADER_FILL
    c.border = BORDER
ws.row_dimensions[HR].height = 24

# 4) Du lieu
rows = [
    [1, "07/10/2026", "Dựng lab phân chia VLAN trên Packet Tracer", "Trần Bách Thăng", 4, "Hoàn thành", ""],
    [2, "07/10/2026", "Điền form báo cáo tuần QTC", "Trần Bách Thăng", 2, "Đang làm", "chờ anh minh chứng"],
    [3, "08/10/2026", "Ôn tập HTB - web attacks", "Trần Bách Thăng", 3, "Chưa làm", ""],
    [4, "10-Sep", "khong co gi ca ne", "TBT", 4, "huhu", "k bic meowmeow"],
]
for i, r in enumerate(rows):
    rr = HR + 1 + i
    for j, v in enumerate(r, 1):
        c = ws.cell(row=rr, column=j, value=v)
        c.border = BORDER
        c.alignment = CENTER if j in (1, 2, 5, 6) else LEFT
    ws.row_dimensions[rr].height = 24

# 5) Do rong cot
widths = [6, 14, 42, 18, 9, 14, 22]
for j, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(j)].width = w

ws.freeze_panes = "A4"

wb.save(OUT)
print("OK ->", OUT)
