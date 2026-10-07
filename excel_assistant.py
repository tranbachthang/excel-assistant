#!/usr/bin/env python3
"""excel-assistant - tro ly Excel nhe: doc / tao / dien form / format / toi uu file .xlsx.

Chi phu thuoc openpyxl. Khong can Excel cai san.

Vi du:
  python excel_assistant.py info    book.xlsx
  python excel_assistant.py read    book.xlsx --head 10
  python excel_assistant.py csv2xlsx data.csv out.xlsx
  python excel_assistant.py xlsx2csv book.xlsx out.csv
  python excel_assistant.py beautify book.xlsx
  python excel_assistant.py fill    template.xlsx data.json filled.xlsx
  python excel_assistant.py optimize big.xlsx small.xlsx
  python excel_assistant.py demo
"""
import argparse
import csv
import json
import os
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")


def _open(path):
    if not os.path.isfile(path):
        sys.exit(f"khong thay file: {path}")
    return load_workbook(path)


def cmd_info(a):
    wb = _open(a.file)
    print(f"file   : {a.file} ({os.path.getsize(a.file):,} bytes)")
    print(f"sheets : {len(wb.sheetnames)} -> {', '.join(wb.sheetnames)}")
    for ws in wb.worksheets:
        print(f"  - {ws.title}: {ws.max_row} hang x {ws.max_column} cot")


def cmd_read(a):
    wb = _open(a.file)
    ws = wb[a.sheet] if a.sheet else wb.active
    print(f"[{ws.title}] {ws.max_row} hang")
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if a.head and i > a.head:
            break
        print(" | ".join("" if c is None else str(c) for c in row))


def cmd_csv2xlsx(a):
    wb = Workbook()
    ws = wb.active
    with open(a.csv, newline="", encoding="utf-8-sig") as f:
        for row in csv.reader(f):
            ws.append(row)
    wb.save(a.out)
    print(f"OK {a.out} ({ws.max_row} hang)")


def cmd_xlsx2csv(a):
    wb = _open(a.xlsx)
    ws = wb[a.sheet] if a.sheet else wb.active
    with open(a.out, "w", newline="", encoding="utf-8-sig") as f:
        csv.writer(f).writerows(ws.iter_rows(values_only=True))
    print(f"OK {a.out}")


def _beautify_sheet(ws):
    if ws.max_row < 1:
        return
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                c.border = BORDER
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        width = max((len(str(c.value)) for c in ws[letter] if c.value is not None), default=8)
        ws.column_dimensions[letter].width = min(max(width + 2, 8), 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def cmd_beautify(a):
    wb = _open(a.file)
    for ws in wb.worksheets:
        _beautify_sheet(ws)
    wb.save(a.file)
    print(f"OK da format {a.file} ({len(wb.sheetnames)} sheet)")


def cmd_fill(a):
    """Dien gia tri vao template theo dia chi o. data.json = {"Sheet1": {"B2": "x", ...}}"""
    wb = _open(a.template)
    data = json.load(open(a.data, encoding="utf-8"))
    n = 0
    for sheet, cells in data.items():
        ws = wb[sheet] if sheet in wb.sheetnames else wb.active
        for addr, val in cells.items():
            ws[addr] = val
            n += 1
    out = a.out or a.template
    wb.save(out)
    print(f"OK dien {n} o -> {out}")


def cmd_optimize(a):
    wb = _open(a.file)
    wb.save(a.out)  # openpyxl ghi lai se bo style/part thua
    before, after = os.path.getsize(a.file), os.path.getsize(a.out)
    print(f"OK {before:,} -> {after:,} bytes ({100 * (before - after) // max(before, 1)}% giam)")


def cmd_demo(_a):
    import tempfile
    tmp = tempfile.mkdtemp()
    form = os.path.join(tmp, "form.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Phieu"
    ws.append(["Ho ten", "Ngay", "Cong viec", "So gio"])
    for _ in range(3):
        ws.append([None, None, None, None])
    wb.save(form)

    data = os.path.join(tmp, "data.json")
    json.dump({"Phieu": {"A2": "Tran Bach Thang", "B2": "07/10/2026", "C2": "Dung lab VLAN", "D2": 4}},
              open(data, "w", encoding="utf-8"), ensure_ascii=False)

    out = os.path.join(tmp, "filled.xlsx")
    cmd_fill(argparse.Namespace(template=form, data=data, out=out))
    ws2 = load_workbook(out).active
    assert ws2["A2"].value == "Tran Bach Thang", ws2["A2"].value
    assert ws2["D2"].value == 4, ws2["D2"].value
    print("SELFTEST PASS")


def main():
    p = argparse.ArgumentParser(prog="excel-assistant", description="tro ly Excel nhe (openpyxl)")
    sub = p.add_subparsers(dest="cmd", required=True)

    s = sub.add_parser("info"); s.add_argument("file"); s.set_defaults(fn=cmd_info)
    s = sub.add_parser("read"); s.add_argument("file"); s.add_argument("--sheet"); s.add_argument("--head", type=int, default=0); s.set_defaults(fn=cmd_read)
    s = sub.add_parser("csv2xlsx"); s.add_argument("csv"); s.add_argument("out"); s.set_defaults(fn=cmd_csv2xlsx)
    s = sub.add_parser("xlsx2csv"); s.add_argument("xlsx"); s.add_argument("out"); s.add_argument("--sheet"); s.set_defaults(fn=cmd_xlsx2csv)
    s = sub.add_parser("beautify"); s.add_argument("file"); s.set_defaults(fn=cmd_beautify)
    s = sub.add_parser("fill"); s.add_argument("template"); s.add_argument("data"); s.add_argument("out", nargs="?"); s.set_defaults(fn=cmd_fill)
    s = sub.add_parser("optimize"); s.add_argument("file"); s.add_argument("out"); s.set_defaults(fn=cmd_optimize)
    s = sub.add_parser("demo"); s.set_defaults(fn=cmd_demo)

    a = p.parse_args()
    a.fn(a)


if __name__ == "__main__":
    main()
