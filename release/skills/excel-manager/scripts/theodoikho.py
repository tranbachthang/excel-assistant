#!/usr/bin/env python3
"""theodoikho.py - Tu dong tao sheet theo doi kho cho cac ngay tiep theo.

Tu mau sheet "01-09", tao them cac sheet ngay (02-09, 03-09, ...).
Moi sheet ngay:
  - Cap nhat o Ngay (B2) va Thu (C2).
  - Cot "Ton dau ngay" (E5:E19) tham chieu "Ton cuoi ngay" (cot H) cua sheet ngay TRUOC.
  - Cot Nhap/Xuat/Ghi chu de trong de nguoi dung nhap.

Dung:
  python theodoikho.py <file.xlsx> [--den DD/MM/YYYY] [--out out.xlsx]

  --den  : tao cac sheet ngay tu (ngay ke sheet cuoi) den ngay nay.
           Bo qua -> chi THEM 1 ngay (ngay ke tiep).
  --out  : file xuat (mac dinh ghi de vao file goc).

Toi uu: lazy-import openpyxl (--help nhanh).
"""
import argparse
import os
import re
import sys

# Unicode tieng Viet tren Windows (tranh loi khi xuat text ra console)
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

from datetime import date, datetime, timedelta

THU_FULL = ["Thứ Hai", "Thứ Ba", "Thứ Tư", "Thứ Năm", "Thứ Sáu", "Thứ Bảy", "Chủ Nhật"]
SHEET_DAY = re.compile(r"^(\d{2})-(\d{2})$")
MAU = "01-09"          # sheet mau goc
ROW_START, ROW_END = 5, 19   # vung du lieu hang hoa


def parse_date_arg(s):
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(s.strip(), fmt).date()
        except ValueError:
            pass
    sys.exit(f"[!] Ngay khong hop le: {s} (dung DD/MM/YYYY)")


def sheet_date(ws, year):
    """Tra ve datetime.date neu sheet co ten dang DD-MM, nguoc lai None."""
    m = SHEET_DAY.match(ws.title.strip())
    if not m:
        return None
    d, mo = int(m.group(1)), int(m.group(2))
    try:
        return date(year, mo, d)
    except ValueError:
        return None


def set_cell(ws, coord, value):
    """Ghi gia tri vao o, xu ly ca o nam trong vung merge (ghi vao top-left)."""
    cell = ws[coord]
    for rng in ws.merged_cells.ranges:
        if cell.coordinate in rng:
            ws.cell(row=rng.min_row, column=rng.min_col).value = value
            return
    cell.value = value


def tao_sheet(wb, mau_ws, new_date, prev_title):
    """Copy sheet mau -> sheet ngay moi, cap nhat cong thuc."""
    dst = wb.copy_worksheet(mau_ws)
    dst.title = f"{new_date.day:02d}-{new_date.month:02d}"

    set_cell(dst, "B2", new_date)
    set_cell(dst, "C2", THU_FULL[new_date.weekday()])

    for r in range(ROW_START, ROW_END + 1):
        dst.cell(row=r, column=5).value = f"=IF(B{r}=\"\",\"\",'{prev_title}'!H{r})"
        dst.cell(row=r, column=6).value = None
        dst.cell(row=r, column=7).value = None
        dst.cell(row=r, column=9).value = None

    return dst


def main():
    p = argparse.ArgumentParser(description="Tao sheet theo doi kho cho cac ngay tiep theo")
    p.add_argument("file", help="File .xlsx theo doi kho (co sheet mau 01-09)")
    p.add_argument("--den", help="Ngay cuoi cung can tao (DD/MM/YYYY). Bo qua -> them 1 ngay.")
    p.add_argument("--out", help="File xuat (mac dinh ghi de file goc)")
    a = p.parse_args()

    if not os.path.isfile(a.file):
        sys.exit(f"[!] Khong thay file: {a.file}")

    from openpyxl import load_workbook  # lazy import

    wb = load_workbook(a.file)
    if MAU not in wb.sheetnames:
        sys.exit(f"[!] Khong thay sheet mau '{MAU}' trong file.")

    mau_ws = wb[MAU]
    ngay_mau = mau_ws["B2"].value
    if isinstance(ngay_mau, datetime):
        year = ngay_mau.year
    elif isinstance(ngay_mau, date):
        year = ngay_mau.year
    else:
        year = parse_date_arg(str(ngay_mau)).year if ngay_mau else datetime.now().year

    existing = []
    for ws in wb.worksheets:
        d = sheet_date(ws, year)
        if d is not None:
            existing.append((d, ws.title))
    existing.sort()
    if not existing:
        sys.exit("[!] Khong thay sheet ngay nao (dang DD-MM) de xac dinh ngay bat dau.")
    last_date, last_title = existing[-1]

    den = parse_date_arg(a.den) if a.den else last_date + timedelta(days=1)
    if den <= last_date:
        sys.exit(f"[!] Ngay den ({den}) phai lon hon ngay cuoi hien co ({last_date}).")

    n = 0
    prev_title = last_title
    cur = last_date + timedelta(days=1)
    while cur <= den:
        new_title = f"{cur.day:02d}-{cur.month:02d}"
        tao_sheet(wb, mau_ws, cur, prev_title)
        print(f"[*] Da tao sheet {new_title} (ton dau tham chieu '{prev_title}')")
        prev_title = new_title
        n += 1
        cur += timedelta(days=1)

    out = a.out or a.file
    wb.save(out)
    print(f"OK -> {os.path.abspath(out)} ({n} sheet ngay moi)")


if __name__ == "__main__":
    main()
