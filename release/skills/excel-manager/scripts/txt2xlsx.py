#!/usr/bin/env python3
"""txt2xlsx.py - Chuyen du lieu text (cot cach boi khoang trang) -> file Excel co format.

Dung cho du lieu co cau truc: STT | Ngay | Cong viec | Nguoi thuc hien | So gio | Trang thai | Ghi chu
Parse bang regex voi danh sach gia tri huu han (chinh xac 100%).
"""
import re
import sys

# Unicode tieng Viet tren Windows (tranh loi khi xuat text ra console)
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

CONGVIEC = (
    r"Vá lỗ hổng CVE|Viết báo cáo pentest|Cấu hình VLAN/switch|Dựng lab Packet Tracer|"
    r"Đánh giá cấu hình router|Rà soát log firewall|Cập nhật tài liệu HTB|"
    r"Họp nhóm kỹ thuật|Kiểm thử xâm nhập web|Kiểm tra backup|Phân tích mã độc|"
    r"Sửa script tự động"
)
NGUOI = (
    r"Trần Bách Thăng|Nguyễn Văn An|Võ Thu Dung|Đặng Hoàng Em|Phạm Minh Cường|Lê Thị Bình"
)
TRANGTHAI = r"Chưa làm|Hoàn thành|Đang làm|Tạm dừng"
GHICHU = r"chờ review|ưu tiên"

PATTERN = re.compile(
    r"^\s*(\d+)\s+(\d{2}/\d{2}/\d{4})\s+(" + CONGVIEC + r")\s+(" + NGUOI +
    r")\s+(\d)\s+(" + TRANGTHAI + r")(?:\s+(" + GHICHU + r"))?\s*$"
)


def parse_line(line):
    m = PATTERN.match(line)
    if not m:
        return None
    stt, ngay, cv, nguoi, gio, tt = m.group(1), m.group(2), m.group(3), m.group(4), int(m.group(5)), m.group(6)
    ghichu = m.group(7) or ""
    return [stt, ngay, cv, nguoi, gio, tt, ghichu]


def build(src, out):
    rows = []
    errors = []
    with open(src, encoding="utf-8") as f:
        for i, line in enumerate(f, 1):
            line = line.rstrip("\n")
            if not line.strip():
                continue
            parsed = parse_line(line)
            if parsed is None:
                errors.append(i)
            else:
                rows.append(parsed)

    wb = Workbook()
    ws = wb.active
    ws.title = "DuLieuCongViec"

    # Tieu de
    NCOL = 7
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NCOL)
    c = ws.cell(row=1, column=1, value=f"BẢNG DỮ LIỆU CÔNG VIỆC ({len(rows)} DÒNG)")
    c.font = Font(bold=True, size=16)
    c.alignment = CENTER
    ws.row_dimensions[1].height = 30

    # Header
    headers = ["STT", "Ngày", "Công việc", "Người thực hiện", "Số giờ", "Trạng thái", "Ghi chú"]
    HR = 2
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=HR, column=j, value=h)
        c.font = Font(bold=True)
        c.alignment = CENTER
        c.fill = HEADER_FILL
        c.border = BORDER
    ws.row_dimensions[HR].height = 24

    # Du lieu
    for i, r in enumerate(rows):
        rr = HR + 1 + i
        for j, v in enumerate(r, 1):
            c = ws.cell(row=rr, column=j, value=v)
            c.border = BORDER
            c.alignment = CENTER if j in (1, 2, 5, 6) else LEFT
    ws.freeze_panes = "A3"

    # Do rong cot
    widths = [8, 14, 34, 20, 10, 14, 16]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w

    # Auto filter
    ws.auto_filter.ref = f"A{HR}:{get_column_letter(NCOL)}{HR + len(rows)}"

    wb.save(out)
    print(f"OK -> {out}")
    print(f"[*] Da parse {len(rows)} dong thanh cong.")
    if errors:
        print(f"[!] {len(errors)} dong LOI (khong parse duoc): {errors[:20]}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python txt2xlsx.py <data.txt> [out.xlsx]")
        sys.exit(1)
    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else "output.xlsx"
    build(src, out)
